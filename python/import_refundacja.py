"""
Imports Polish Ministry of Health (MZ/NFZ) Drug Reimbursement List (Wykaz leków refundowanych)
from Excel (.xlsx) into SQLite database `data/rpl.db`.
"""

import os
import sys
import sqlite3
import argparse
import openpyxl
from rich.console import Console
from rich.table import Table

console = Console()

DEFAULT_XLSX_PATH = "/home/karol/Pobrane/83W_zalacznik_do_obwieszczenia.xlsx"
DEFAULT_DB_PATH = "data/rpl.db"


def normalize_gtin(gtin) -> str:
    if not gtin:
        return ""
    g_str = str(gtin).strip().replace(" ", "").replace("\t", "").replace("\n", "")
    return g_str.lstrip("0")


def init_refundacja_table(conn: sqlite3.Connection):
    conn.execute("DROP TABLE IF EXISTS refundacja;")
    conn.execute("""
    CREATE TABLE refundacja (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        kod_gtin TEXT NOT NULL,
        kod_gtin_norm TEXT NOT NULL,
        typ_listy TEXT NOT NULL, -- 'A1', 'A2', 'A3', 'B', 'C', 'D1', 'D2', 'E'
        substancja_czynna TEXT,
        nazwa_lek_dawka TEXT,
        zawartosc_opakowania TEXT,
        grupa_limitowa TEXT,
        cena_zbytu_netto TEXT,
        urzedowa_cena_zbytu TEXT,
        cena_hurtowa_brutto TEXT,
        cena_detaliczna TEXT,
        wysokosc_limitu TEXT,
        zakres_wskazan TEXT,
        zakres_wskazan_pozarejestracyjnych TEXT,
        poziom_odplatnosci TEXT,
        wysokosc_doplaty TEXT,
        bezplatny_dziecko_18 INTEGER DEFAULT 0,
        bezplatny_senior_65 INTEGER DEFAULT 0,
        bezplatny_ciaza INTEGER DEFAULT 0
    );
    """)
    conn.execute("CREATE INDEX idx_refundacja_gtin_norm ON refundacja(kod_gtin_norm);")
    conn.execute("CREATE INDEX idx_refundacja_typ ON refundacja(typ_listy);")
    conn.commit()


def parse_and_import_refundacja(xlsx_path: str, db_path: str):
    console.print(f"[cyan]Loading Excel workbook: {xlsx_path}...[/cyan]")
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    
    conn = sqlite3.connect(db_path)
    init_refundacja_table(conn)

    # 1. First collect special flags from D1 (dzieci < 18), D2 (seniorzy 65+), E (kobiety w ciąży)
    gtins_d1_dzieci = set()
    if "D1" in wb.sheetnames:
        for i, row in enumerate(wb["D1"].iter_rows(values_only=True)):
            if i >= 2 and len(row) > 4 and row[4]:
                gtins_d1_dzieci.add(normalize_gtin(row[4]))

    gtins_d2_seniorzy = set()
    if "D2" in wb.sheetnames:
        for i, row in enumerate(wb["D2"].iter_rows(values_only=True)):
            if i >= 2 and len(row) > 4 and row[4]:
                gtins_d2_seniorzy.add(normalize_gtin(row[4]))

    gtins_e_ciaza = set()
    if "E" in wb.sheetnames:
        for i, row in enumerate(wb["E"].iter_rows(values_only=True)):
            if i >= 2 and len(row) > 4 and row[4]:
                gtins_e_ciaza.add(normalize_gtin(row[4]))

    console.print(f"• Flagi specjalne: <18 ({len(gtins_d1_dzieci)} GTIN), 65+ ({len(gtins_d2_seniorzy)} GTIN), Ciąża+ ({len(gtins_e_ciaza)} GTIN)")

    # 2. Parse main reimbursement sheets: A1, A2, A3, B, C
    total_imported = 0
    records_to_insert = []

    for sheet_name in ["A1", "A2", "A3", "B", "C"]:
        if sheet_name not in wb.sheetnames:
            continue

        sheet = wb[sheet_name]
        console.print(f"[cyan]Parsing sheet {sheet_name}...[/cyan]")
        count_sheet = 0

        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i < 2:  # Skip title and header
                continue
            if not row or len(row) < 5 or not row[4]:
                continue

            raw_gtin = str(row[4]).strip()
            norm_gtin = normalize_gtin(raw_gtin)
            if not norm_gtin:
                continue

            subst = str(row[1] or "").strip()
            nazwa = str(row[2] or "").strip()
            opak = str(row[3] or "").strip()
            grupa_lim = str(row[7] or "").strip() if len(row) > 7 else ""
            cena_zbytu_netto = str(row[8] or "").strip() if len(row) > 8 else ""
            urzedowa_cena = str(row[9] or "").strip() if len(row) > 9 else ""
            cena_hurt_brutto = str(row[10] or "").strip() if len(row) > 10 else ""
            cena_detaliczna = str(row[11] or "").strip() if len(row) > 11 else ""
            limit = str(row[12] or "").strip() if len(row) > 12 else ""
            wskazania = str(row[13] or "").strip() if len(row) > 13 else ""
            wskazania_pozarej = str(row[14] or "").strip() if len(row) > 14 else ""
            poziom_odp = str(row[15] or "").strip() if len(row) > 15 else ""
            doplata = str(row[16] or "").strip() if len(row) > 16 else ""

            # For sheets B (programy lekowe) and C (chemioterapia), co-payment is 100% covered by NFZ (bezpłatny)
            if sheet_name in ["B", "C"]:
                poziom_odp = "bezpłatny (program lekowy / chemioterapia)"
                doplata = "0,00"

            is_dziecko = 1 if norm_gtin in gtins_d1_dzieci else 0
            is_senior = 1 if norm_gtin in gtins_d2_seniorzy else 0
            is_ciaza = 1 if norm_gtin in gtins_e_ciaza else 0

            records_to_insert.append((
                raw_gtin,
                norm_gtin,
                sheet_name,
                subst,
                nazwa,
                opak,
                grupa_lim,
                cena_zbytu_netto,
                urzedowa_cena,
                cena_hurt_brutto,
                cena_detaliczna,
                limit,
                wskazania,
                wskazania_pozarej,
                poziom_odp,
                doplata,
                is_dziecko,
                is_senior,
                is_ciaza
            ))
            count_sheet += 1

        total_imported += count_sheet
        console.print(f"[green]  &bull; {sheet_name}: zaimportowano {count_sheet} wpisów.[/green]")

    # Batch insert
    conn.executemany("""
    INSERT INTO refundacja (
        kod_gtin, kod_gtin_norm, typ_listy, substancja_czynna, nazwa_lek_dawka,
        zawartosc_opakowania, grupa_limitowa, cena_zbytu_netto, urzedowa_cena_zbytu,
        cena_hurtowa_brutto, cena_detaliczna, wysokosc_limitu, zakres_wskazan,
        zakres_wskazan_pozarejestracyjnych, poziom_odplatnosci, wysokosc_doplaty,
        bezplatny_dziecko_18, bezplatny_senior_65, bezplatny_ciaza
    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, records_to_insert)
    conn.commit()

    # Match check against RPL database opakowania
    c = conn.cursor()
    c.execute("""
    SELECT count(DISTINCT o.produkt_id)
    FROM opakowania o
    JOIN refundacja r ON ltrim(o.kod_gtin, '0') = r.kod_gtin_norm
    """)
    matched_products = c.fetchone()[0]

    c.execute("SELECT count(*) FROM refundacja")
    total_refund_rows = c.fetchone()[0]

    conn.close()

    console.print(f"\n[bold green]Sukces: Zaimportowano {total_refund_rows:,} pozycji refundacyjnych![/bold green]")
    console.print(f"• Powiązanych produktów leczniczych w bazie RPL: [bold yellow]{matched_products:,}[/bold yellow]\n")


def main():
    parser = argparse.ArgumentParser(description="Import MZ/NFZ Drug Reimbursement Excel List into SQLite")
    parser.add_argument("--xlsx", type=str, default=DEFAULT_XLSX_PATH, help="Path to obwieszczenie .xlsx file")
    parser.add_argument("--db", type=str, default=DEFAULT_DB_PATH, help="Path to SQLite database")
    args = parser.parse_args()

    parse_and_import_refundacja(args.xlsx, args.db)


if __name__ == "__main__":
    main()
